import tempfile
import unittest
import io
import zipfile
from pathlib import Path
from unittest.mock import patch

import app
from tools import platform_sync


class CalculationTests(unittest.TestCase):
    def test_connector_python_uses_current_interpreter_outside_local_venv(self):
        with patch.dict("app.os.environ", {"WORKCOUNT_CONNECTOR_PYTHON": "/usr/local/bin/python"}):
            self.assertEqual(app.connector_python(), "/usr/local/bin/python")

    def test_frozen_connector_is_next_to_server_bundle(self):
        with patch.object(app.sys, "frozen", True, create=True), patch.object(app, "ROOT", Path("/tmp/workcount-portable")), patch.object(app.os, "name", "nt"):
            self.assertEqual(app.connector_command(), ["/tmp/workcount-portable/PlatformConnector/PlatformConnector.exe"])

    def test_connector_environment_uses_native_temp_and_utf8(self):
        environment = app.connector_environment()
        self.assertEqual(environment["PYTHONIOENCODING"], "utf-8")
        self.assertEqual(environment["PYTHONUTF8"], "1")
        self.assertTrue(environment["PYTHONPYCACHEPREFIX"].startswith(tempfile.gettempdir()))

    def test_connector_result_preserves_startup_error(self):
        with self.assertRaisesRegex(app.AppError, "chromedriver failed"):
            app.connector_result("", "chromedriver failed", 1, "平台登录组件未能正常启动")

    def test_windows_browser_candidates_prefer_chrome_and_fall_back_to_edge(self):
        environment = {"PROGRAMFILES": "C:/Program Files", "PROGRAMFILES(X86)": "C:/Program Files (x86)"}
        existing = {
            "C:/Program Files/Google/Chrome/Application/chrome.exe",
            "C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe",
        }
        candidates = platform_sync.windows_browser_candidates(environment, existing.__contains__)
        self.assertEqual(candidates, [
            ("Chrome", "C:/Program Files/Google/Chrome/Application/chrome.exe"),
            ("Edge", "C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe"),
        ])

    def test_qingguo_session_cookie_is_restored_without_second_login(self):
        class Driver:
            current_url = ""
            cookies = []

            def get(self, url):
                self.current_url = url if not self.cookies else platform_sync.BASE + "/frame/homes.action"

            def add_cookie(self, cookie):
                self.cookies.append(cookie)

        driver = Driver()
        restored = platform_sync.restore_session(driver, [{"name": "JSESSIONID", "value": "secret", "unexpected": "drop"}])
        self.assertTrue(restored)
        self.assertEqual(driver.cookies, [{"name": "JSESSIONID", "value": "secret"}])

    def test_qingguo_session_captures_cas_path_cookie(self):
        class Driver:
            def execute_cdp_cmd(self, command, params):
                self.command = (command, params)
                return {"cookies": [
                    {"name": "CASTGC", "value": "ticket", "domain": "qgjw.suet.edu.cn", "path": "/cas"},
                    {"name": "other", "value": "drop", "domain": "example.com", "path": "/"},
                ]}

        driver = Driver()
        self.assertEqual(platform_sync.session_cookies(driver), [
            {"name": "CASTGC", "value": "ticket", "domain": "qgjw.suet.edu.cn", "path": "/cas"},
        ])
        self.assertEqual(driver.command, ("Network.getAllCookies", {}))

    def test_login_returns_qingguo_error_instead_of_generic_message(self):
        class Element:
            text = "账号或密码有误!"

            def click(self): pass
            def send_keys(self, value): pass

        class Driver:
            current_url = platform_sync.BASE + "/cas/login.action"

            def get(self, url): self.current_url = url
            def find_element(self, by, value): return Element()
            def execute_script(self, script): return None

        class Wait:
            def __init__(self, driver, timeout): self.driver = driver
            def until(self, predicate): return predicate(self.driver)

        class Conditions:
            @staticmethod
            def presence_of_element_located(locator): return lambda driver: True

        class Locator:
            ID = "id"

        with patch.object(platform_sync, "WebDriverWait", Wait), patch.object(platform_sync, "EC", Conditions), patch.object(platform_sync, "By", Locator):
            with self.assertRaisesRegex(RuntimeError, "青果平台提示：账号或密码有误"):
                platform_sync.login(Driver(), "user", "bad-password")

    def test_login_accepts_session_when_qingguo_redirect_is_slow(self):
        class Element:
            text = "正在登录......"

            def click(self): pass
            def send_keys(self, value): pass

        class Driver:
            current_url = platform_sync.BASE + "/cas/login.action"
            login_result = None

            def get(self, url):
                if url.endswith("/frame/homes.action"):
                    self.current_url = url
                else:
                    self.current_url = platform_sync.BASE + "/cas/login.action"

            def find_element(self, by, value):
                return Element()

            def execute_script(self, script):
                return self.login_result

        class Wait:
            calls = 0

            def __init__(self, driver, timeout): self.driver = driver
            def until(self, predicate):
                Wait.calls += 1
                if Wait.calls == 2:
                    raise platform_sync.TimeoutException("slow JavaScript redirect")
                return predicate(self.driver)

        class Conditions:
            @staticmethod
            def presence_of_element_located(locator): return lambda driver: True

        class Locator:
            ID = "id"

        with patch.object(platform_sync, "WebDriverWait", Wait), patch.object(platform_sync, "EC", Conditions), patch.object(platform_sync, "By", Locator):
            platform_sync.login(Driver(), "user", "password")
        self.assertEqual(Wait.calls, 3)

    def test_login_retries_once_after_ajax_network_error(self):
        class Element:
            text = "正在登录......"
            def click(self): pass
            def send_keys(self, value): pass

        class Driver:
            current_url = platform_sync.BASE + "/cas/login.action"
            attempts = 0
            def get(self, url):
                self.current_url = url
                if url.endswith("/cas/login.action"):
                    self.attempts += 1
            def find_element(self, by, value): return Element()
            def execute_script(self, script):
                if script.startswith("return"):
                    return {"state": "network_error"} if self.attempts == 1 else {"state": "response", "data": {"status": "200"}}

        class Wait:
            def __init__(self, driver, timeout): self.driver = driver; self.timeout = timeout
            def until(self, predicate):
                if self.timeout == platform_sync.SESSION_PROBE_TIMEOUT:
                    self.driver.current_url = platform_sync.BASE + "/cas/login.action"
                    return False
                result = predicate(self.driver)
                if self.timeout == platform_sync.LOGIN_RESPONSE_TIMEOUT and self.driver.attempts == 2:
                    self.driver.current_url = platform_sync.BASE + "/frame/homes.action"
                return result

        class Conditions:
            @staticmethod
            def presence_of_element_located(locator): return lambda driver: True

        class Locator:
            ID = "id"

        with patch.object(platform_sync, "WebDriverWait", Wait), patch.object(platform_sync, "EC", Conditions), patch.object(platform_sync, "By", Locator):
            platform_sync.login(Driver(), "user", "password")

    def test_production_cookie_requires_https(self):
        with patch.object(app, "SECURE_COOKIES", True):
            self.assertIn("; Secure", app.session_cookie("token", 60))
        with patch.object(app, "SECURE_COOKIES", False):
            self.assertNotIn("; Secure", app.session_cookie("token", 60))

    def test_export_templates_are_present(self):
        for name in app.REQUIRED_TEMPLATES:
            self.assertTrue((app.TEMPLATES / name).is_file(), name)

    def test_theory_matches_source_workbook(self):
        row = {"kind":"theory","course":"Python高级应用","student_count":60,"total_hours":46,"experiment_hours":18,"experiment_students":60,"category_coeff":1.2,"repeat_coeff":1,"course_coeff":1,"online_coeff":1}
        self.assertEqual(app.calculate(row)["workload"], 62.88)

    def test_complete_seed_matches_source_workbook(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(app, "DB_PATH", Path(tmp) / "test.db"):
            app.init_db(seed_demo=True); items = app.item_list(); totals = app.summary(items)
            self.assertEqual(totals["theory_hours"], 222)
            self.assertEqual(totals["theory_workload"], 257.6992)
            self.assertEqual(totals["practice_hours"], 87)
            self.assertEqual(totals["practice_workload"], 108.15)
            self.assertEqual(totals["total_workload"], 365.8492)

    def test_validation_rejects_unknown_kind(self):
        with self.assertRaises(app.AppError): app.validate_item({"kind":"other","course":"x"})

    def test_overview_practice_hours_match_workbook_rules(self):
        internship = app.calculate({"kind":"internship", "student_count":30, "total_hours":10, "weeks":10, "practice_coeff":4, "instructors":1})
        thesis = app.calculate({"kind":"thesis", "student_count":6, "total_hours":10, "weeks":10, "practice_coeff":6, "instructors":1})
        self.assertEqual(internship["practice_hours"], 150)
        self.assertEqual(thesis["practice_hours"], 36)
        self.assertEqual(internship["display_hours"], 150)

    def test_course_experiment_workload_stays_in_theory_course_total(self):
        item = app.calculate({"kind":"theory", "course":"课程A", "student_count":50, "total_hours":32, "experiment_hours":16, "experiment_students":50, "category_coeff":1, "repeat_coeff":1, "course_coeff":1, "online_coeff":1})
        self.assertEqual(item["workload"], 32)
        self.assertEqual(item["theory_workload"], 32)
        self.assertEqual(item["practice_workload"], 0)

    def test_formula_export_is_auditable(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(app, "DB_PATH", Path(tmp) / "test.db"):
            app.init_db(seed_demo=True)
            raw = app.make_formula_xlsx(app.item_list("0000000", "示例教师"))
            self.assertTrue(raw.startswith(b"PK"))
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                self.assertIn(b"SUMIF", archive.read("xl/worksheets/sheet1.xml"))

    def test_confirmed_term_start_dates(self):
        self.assertEqual(platform_sync.first_monday(2025, 1).isoformat(), "2026-03-09")
        self.assertEqual(platform_sync.first_monday(2026, 0).isoformat(), "2026-08-31")

    def test_schedule_events_follow_qingguo_without_inventing_note_periods(self):
        html = """
        <div id="015">课程A 考试 [5]周 1-2节 50 教室 本部（中心校区） 班级A</div>
        <div id="025">课程A 考试 [5]周 1-2节 50 教室 本部（中心校区） 班级A</div>
        <div>注1：[X]集中实训、[6周]、班级A(第1批50人)</div>
        """
        events = platform_sync.parse_schedule_events(html, 2025, 1, "1", "教师")
        self.assertEqual(len(events), 2)
        self.assertEqual([event["event_date"] for event in events], ["2026-04-06", "2026-04-07"])
        self.assertFalse(any(event["periods"] == "1-6" for event in events))

    def test_personal_hours_combined_weeks_keep_one_label_and_sum_both(self):
        from openpyxl import load_workbook

        months = [
            {"semester":"2025-2026-2", "employee_id":"1", "employee_name":"教师", "college":"人工智能学院", "department":"大数据教研室", "month_key":month, "theory_hours":0, "practice_hours":0}
            for month in ("3月", "4月", "5月", "6-7月")
        ]
        events = [
            {"semester":"2025-2026-2", "employee_id":"1", "employee_name":"教师", "event_date":day, "academic_week":week, "weekday":1, "course":"课程A", "class_name":"班级A", "student_count":50, "periods":"5-6", "hours":2, "category":"practice", "source":"平台:教学安排"}
            for week, day in ((17, "2026-06-29"), (18, "2026-07-06"))
        ]
        raw = app.exact_personal_hours_xlsx("1", "教师", "2025-2026-2", months, events)
        workbook = load_workbook(io.BytesIO(raw), data_only=False)
        sheet = workbook["6-7月"]
        self.assertEqual(sheet["AA14"].value, "5-6")
        self.assertEqual(sheet["AF14"].value, "=2+2")
        self.assertEqual(sheet["AF3"].value, "第十七周、第十八周")
        self.assertEqual(sheet["AA3"].value, sheet["AA12"].value)
        self.assertEqual(sheet["AF3"].value, sheet["AF12"].value)

    def test_fall_personal_hours_removes_template_holiday_notes(self):
        from openpyxl import load_workbook

        months = [
            {"semester":"2024-2025-1", "employee_id":"1", "employee_name":"教师", "college":"人工智能学院", "department":"大数据教研室", "month_key":month, "theory_hours":0, "practice_hours":0}
            for month in ("9月", "10月", "11月", "12-1月")
        ]
        events = [{"semester":"2024-2025-1", "employee_id":"1", "employee_name":"教师", "event_date":"2024-10-08", "academic_week":6, "weekday":2, "course":"课程A", "class_name":"班级A", "student_count":50, "periods":"1-2", "hours":2, "category":"practice", "source":"平台:教学安排"}]
        raw = app.exact_personal_hours_xlsx("1", "教师", "2024-2025-1", months, events)
        workbook = load_workbook(io.BytesIO(raw), data_only=False)
        for sheet in workbook.worksheets[1:]:
            text = "\n".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
            self.assertNotIn("清明节", text)
            self.assertNotIn("劳动节", text)
            self.assertNotIn("端午节", text)
            self.assertEqual(sheet["A21"].value, "备注：课程、周次和上课日期均依据青果教学安排生成。")

    def test_qingguo_no_practice_record_does_not_invent_workload(self):
        html = "<html><body><div>山东工程职业技术大学教师指导实践环节</div><div>没有检索到记录!</div></body></html>"
        self.assertEqual(platform_sync.parse_practice(html, "2024-2025-1", "1", "教师"), [])

    def test_decision_export_expands_practice_rows_without_truncation(self):
        from openpyxl import load_workbook

        base = {"kind":"training", "employee_id":"1", "employee_name":"教师", "semester":"2025-2026-2", "course_code":"X", "student_count":50, "total_hours":10, "experiment_hours":0, "experiment_students":0, "category_coeff":1, "repeat_coeff":1, "course_coeff":1, "online_coeff":1, "practice_coeff":0.5, "weeks":1, "instructors":1, "enterprise_coeff":1, "manual_workload":0, "source":"test"}
        items = [app.calculate({**base, "course":f"实训{i}", "class_name":f"班级{i}"}) for i in range(7)]
        raw = app.exact_workload_xlsx(items, "1", "教师", "2025-2026-2")
        workbook = load_workbook(io.BytesIO(raw), data_only=False)
        sheet = workbook["实践工作量"]
        self.assertEqual(sheet["H12"].value, "实训6")
        self.assertIn("S6:S12", sheet["AO6"].value)
        self.assertEqual(workbook["总工作量统计表"]["J3"].value, "='实践工作量'!AO6")

    def test_budget_weekly_hours_are_integer_without_changing_platform_total(self):
        from openpyxl import load_workbook

        item = {"kind":"theory", "employee_id":"1", "employee_name":"教师", "semester":"2026-2027-1", "course":"课程A", "class_name":"班级A", "course_code":"X", "student_count":50, "total_hours":50, "experiment_hours":16, "experiment_students":50, "category_coeff":1, "repeat_coeff":1, "course_coeff":1, "online_coeff":1, "practice_coeff":0, "weeks":0, "instructors":1, "enterprise_coeff":1, "manual_workload":0, "source":"test"}
        events = [{"semester":"2026-2027-1", "course":"课程A", "class_name":"班级A", "academic_week":week, "category":"theory"} for week in range(1,18)]
        raw = app.exact_budget_xlsx([item], events, "1", "教师", "2026-2027-1")
        workbook = load_workbook(io.BytesIO(raw), data_only=False)
        sheet = workbook["校内教师理论工作量"]
        self.assertEqual(sheet["Q4"].value, 3)
        self.assertEqual(sheet["R4"].value, 17)
        self.assertEqual(sheet["S4"].value, "=Q4*R4-1")


if __name__ == "__main__": unittest.main()
